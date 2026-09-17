"""Validate artifact structure and complete text extraction; not an ATS score."""
import json
import re
import shutil
import subprocess
import unicodedata
from pathlib import Path
from zipfile import ZipFile

import pdfplumber
from docx import Document
from lxml import etree
from pypdf import PdfReader

from build_resume import DATA, OUT, ROOT, STEM, paragraphs


def normalize(text):
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"(?m)^\s*[-\u2022]\s+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def check():
    qa = ROOT / "qa"
    qa.mkdir(exist_ok=True)
    pdf = OUT / f"{STEM}.pdf"
    docx = OUT / f"{STEM}.docx"
    expected = [text for _, text in paragraphs(DATA)]
    expected_text = normalize("\n".join(expected))
    report = {"scope": "Local artifact and extraction checks; no commercial ATS was run.", "checks": {}}

    def require(name, condition):
        report["checks"][name] = bool(condition)
        if not condition:
            raise AssertionError(name)

    document = Document(docx)
    require("docx_paragraph_content_and_order", [p.text for p in document.paragraphs] == expected)
    require("one_column_no_layout_tables", not document.tables)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(docx) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))
        require("no_text_boxes_or_drawings", not tree.xpath("//w:txbxContent | //w:drawing | //w:pict", namespaces=ns))
        require("no_paragraph_borders", not tree.xpath("//w:pBdr", namespaces=ns) and not styles.xpath("//w:pBdr", namespaces=ns))
        cols = tree.xpath("//w:cols/@w:num", namespaces=ns)
        require("single_column_section", all(c == "1" for c in cols))
        for name in archive.namelist():
            if re.match(r"word/(header|footer)\d+\.xml$", name):
                part = etree.fromstring(archive.read(name))
                require(f"no_content_in_{Path(name).stem}", not part.xpath("//w:t[text()]", namespaces=ns))
    reader = PdfReader(pdf)
    require("pdf_one_letter_page", len(reader.pages) == 1 and abs(float(reader.pages[0].mediabox.width) - 612) < 1 and abs(float(reader.pages[0].mediabox.height) - 792) < 1)
    require("pdf_not_encrypted", not reader.is_encrypted)
    require("docx_below_2_5_MB", docx.stat().st_size < 2_500_000)
    require("pdf_below_2_5_MB", pdf.stat().st_size < 2_500_000)
    extracted = {"pypdf": "\n".join(page.extract_text() for page in reader.pages)}
    pdftotext = shutil.which("pdftotext")
    if not pdftotext:
        raise RuntimeError("pdftotext is required for independent PDF extraction checks")
    for label, args in [("poppler", []), ("poppler_layout", ["-layout"])]:
        extracted[label] = subprocess.check_output([pdftotext, *args, str(pdf), "-"], text=True)
    with pdfplumber.open(pdf) as parsed:
        extracted["pdfplumber"] = "\n".join(page.extract_text() for page in parsed.pages)
        require("text_stays_inside_page_margins", all(
            c["x0"] >= 30 and c["x1"] <= page.width - 30 and c["top"] >= 25 and c["bottom"] <= page.height - 25
            for page in parsed.pages for c in page.chars if c["text"].strip()
        ))
        require("no_raster_images", all(not page.images for page in parsed.pages))
    for label, text in extracted.items():
        (qa / f"{label}.txt").write_text(text)
        require(f"{label}_all_content_in_source_order", normalize(text) == expected_text)
        require(f"{label}_no_replacement_characters", "\ufffd" not in text and "\u0000" not in text)
    report["files"] = {p.name: {"bytes": p.stat().st_size} for p in (docx, pdf)}
    report["paragraphs"] = len(expected)
    report["pages"] = len(reader.pages)
    (qa / "checks.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    check()
